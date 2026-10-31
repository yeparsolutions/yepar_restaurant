# ============================================================
# YEPAR RESTAURANT — backend/servicios/carga_masiva.py
# Carga masiva por Excel (insumos, platos y recetas) y asistente de arranque.
#
# Analogía: es la mudanza. En vez de subir las cajas una por una, el dueño llena
# una planilla y el sistema la revisa entera ANTES de mover nada: o entra todo
# bien, o no entra nada y se le dice en qué fila está cada problema.
#
# Cómo funciona: se aplican las filas dentro de una transacción. "Validar" la
# deshace al final; "Importar" la confirma solo si no hubo ningún error.
# ============================================================

import io
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func
from sqlalchemy.orm import Session

from errores import ErrorNegocio
from models import (Caja, CategoriaItem, ConfiguracionCobro, FamiliaUnidad, IntegracionFacturacion, Item, Mesa,
                    MovimientoStock, Receta, RolUsuario, Sucursal, TipoItem, TipoMovimientoStock, Unidad, Usuario, UsuarioRol)
from servicios.compras import convertir_cantidad, ingresar_compra
from servicios.items import crear_categoria, crear_item
from servicios.recetas import LineaReceta, TIPOS_INGREDIENTE, guardar_receta

MAX_BYTES = 2 * 1024 * 1024
MAX_FILAS = 2000

HOJAS = {
    "insumos": ["nombre", "tipo", "unidad", "categoria", "stock_minimo", "stock_inicial", "costo_por_unidad"],
    "platos": ["nombre", "tipo", "categoria", "precio", "estacion", "stock_inicial", "costo_por_unidad"],
    "recetas": ["plato", "ingrediente", "cantidad", "unidad", "removible", "rendimiento", "unidad_rendimiento"],
}
OBLIGATORIAS = {"insumos": ["nombre", "tipo", "unidad"], "platos": ["nombre", "tipo", "precio"],
                "recetas": ["plato", "ingrediente", "cantidad"]}

TIPOS_INSUMO = {"materia_prima": TipoItem.materia_prima, "insumo": TipoItem.materia_prima,
                "componente": TipoItem.componente, "preparacion": TipoItem.componente}
TIPOS_PLATO = {"preparado": TipoItem.vendible_preparado, "plato": TipoItem.vendible_preparado,
               "reventa": TipoItem.vendible_reventa, "bebida": TipoItem.vendible_reventa}
BASE_POR_FAMILIA = {FamiliaUnidad.peso: "g", FamiliaUnidad.volumen: "ml", FamiliaUnidad.unidad: "und"}


@dataclass
class ErrorFila:
    hoja: str
    fila: int
    mensaje: str


@dataclass
class Reporte:
    errores: list = field(default_factory=list)
    insumos: int = 0
    platos: int = 0
    recetas: int = 0
    con_stock_inicial: int = 0
    categorias_nuevas: int = 0

    @property
    def valido(self) -> bool:
        return not self.errores


# ── Lectura del Excel ────────────────────────────────────────
def _norm(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode().lower().strip()
    return re.sub(r"[^a-z0-9]+", "_", t).strip("_")


def _celda(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def abrir_libro(contenido: bytes):
    if len(contenido) > MAX_BYTES:
        raise ErrorNegocio("El archivo pesa más de 2 MB: divídelo en partes.")
    try:
        return load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise ErrorNegocio("El archivo no es un Excel .xlsx válido (usa la plantilla).")


def _leer_hoja(libro, clave: str, rep: Reporte) -> list[tuple[int, dict]]:
    """Devuelve [(número de fila en Excel, {columna: texto})], saltando filas vacías."""
    hoja = next((libro[n] for n in libro.sheetnames if _norm(n) == clave), None)
    if hoja is None:
        return []
    filas = hoja.iter_rows(values_only=True)
    encabezado = next(filas, None)
    if encabezado is None:
        return []
    cols = [_norm(c) for c in encabezado]
    faltan = [c for c in OBLIGATORIAS[clave] if c not in cols]
    if faltan:
        rep.errores.append(ErrorFila(clave, 1, "Falta la columna: " + ", ".join(faltan) + " (usa la plantilla)."))
        return []
    salida = []
    for n, valores in enumerate(filas, start=2):
        fila = {c: _celda(v) for c, v in zip(cols, valores) if c}
        if not any(fila.values()):
            continue
        if len(salida) >= MAX_FILAS:
            rep.errores.append(ErrorFila(clave, n, f"Máximo {MAX_FILAS} filas por hoja."))
            break
        salida.append((n, fila))
    return salida


# ── Conversión de celdas ─────────────────────────────────────
def _numero(texto: str, campo: str, *, obligatorio=False) -> Optional[Decimal]:
    if texto == "":
        if obligatorio:
            raise ErrorNegocio(f"Falta {campo}.")
        return None
    try:
        n = Decimal(texto.replace(",", ".")) if texto.count(",") == 1 and "." not in texto else Decimal(texto)
    except InvalidOperation:
        raise ErrorNegocio(f"{campo} no es un número ('{texto}').")
    if n < 0:
        raise ErrorNegocio(f"{campo} no puede ser negativo.")
    return n


def _si_no(texto: str) -> bool:
    t = _norm(texto)
    if t in ("", "no", "n", "0", "false"):
        return False
    if t in ("si", "s", "1", "true", "x"):
        return True
    raise ErrorNegocio(f"'{texto}' no es sí/no.")


def _unidad(db: Session, codigo: str) -> Unidad:
    u = db.query(Unidad).filter(func.lower(Unidad.codigo) == codigo.strip().lower()).one_or_none()
    if u is None:
        raise ErrorNegocio(f"La unidad '{codigo}' no existe (usa g, kg, ml, l, und...).")
    return u


def _base_de(db: Session, u: Unidad) -> str:
    return BASE_POR_FAMILIA[u.familia]


class _Contexto:
    """Memoria de la carga: categorías vistas e ítems por nombre (los que ya había y los nuevos)."""

    def __init__(self, db: Session, empresa_id: int, sucursal_id: Optional[int], rep: Reporte):
        self.db, self.empresa_id, self.sucursal_id, self.rep = db, empresa_id, sucursal_id, rep
        self.categorias = {c.nombre.lower(): c.id for c in db.query(CategoriaItem).filter_by(empresa_id=empresa_id)}
        self.items = {i.nombre.lower(): i for i in db.query(Item).filter_by(empresa_id=empresa_id)}
        self.nuevos: set[int] = set()

    def categoria_id(self, nombre: str) -> Optional[int]:
        if not nombre:
            return None
        clave = nombre.lower()
        if clave not in self.categorias:
            self.categorias[clave] = crear_categoria(self.db, empresa_id=self.empresa_id, nombre=nombre).id
            self.rep.categorias_nuevas += 1
        return self.categorias[clave]

    def stock_inicial(self, item: Item, fila: dict, unidad: str):
        cantidad = _numero(fila.get("stock_inicial", ""), "El stock inicial")
        if not cantidad:
            return
        costo = _numero(fila.get("costo_por_unidad", ""), "El costo por unidad")
        if costo is None:
            raise ErrorNegocio("Con stock inicial hay que indicar el costo por unidad (sin IVA).")
        if self.sucursal_id is None:
            raise ErrorNegocio("Para cargar stock inicial indica la sucursal.")
        ingresar_compra(self.db, empresa_id=self.empresa_id, sucursal_id=self.sucursal_id, item_id=item.id,
                        cantidad=cantidad, unidad=unidad, monto_neto=cantidad * costo,
                        referencia_tipo="carga_inicial", nota="Stock inicial (carga masiva)")
        self.rep.con_stock_inicial += 1


# ── Cada hoja ────────────────────────────────────────────────
def _insumo(ctx: _Contexto, f: dict):
    tipo = TIPOS_INSUMO.get(_norm(f["tipo"]))
    if tipo is None:
        raise ErrorNegocio(f"Tipo '{f['tipo']}' no válido (materia_prima o componente).")
    u = _unidad(ctx.db, f["unidad"])
    categoria = ctx.categoria_id(f.get("categoria", ""))
    item = crear_item(ctx.db, empresa_id=ctx.empresa_id, tipo=tipo, nombre=f["nombre"],
                      unidad_base=_base_de(ctx.db, u), categoria_id=categoria)
    minimo = _numero(f.get("stock_minimo", ""), "El stock mínimo")
    if minimo:
        item.stock_minimo = convertir_cantidad(ctx.db, item, cantidad=minimo, unidad=u.codigo)
    if f.get("stock_inicial") and tipo != TipoItem.materia_prima:
        raise ErrorNegocio("Un componente se produce, no se compra: no lleva stock inicial.")
    ctx.stock_inicial(item, f, u.codigo)
    ctx.items[item.nombre.lower()] = item
    ctx.nuevos.add(item.id)
    ctx.rep.insumos += 1


def _plato(ctx: _Contexto, f: dict):
    tipo = TIPOS_PLATO.get(_norm(f["tipo"]))
    if tipo is None:
        raise ErrorNegocio(f"Tipo '{f['tipo']}' no válido (preparado o reventa).")
    precio = _numero(f["precio"], "El precio", obligatorio=True)
    estacion = _norm(f.get("estacion", "")) or ("barra" if tipo == TipoItem.vendible_reventa else "cocina")
    if estacion not in ("cocina", "barra"):
        raise ErrorNegocio(f"Estación '{f['estacion']}' no válida (cocina o barra).")
    if f.get("stock_inicial") and tipo != TipoItem.vendible_reventa:
        raise ErrorNegocio("Un plato preparado no tiene stock propio: su stock son sus ingredientes.")
    item = crear_item(ctx.db, empresa_id=ctx.empresa_id, tipo=tipo, nombre=f["nombre"], unidad_base="und",
                      precio_venta=precio, estacion=estacion, categoria_id=ctx.categoria_id(f.get("categoria", "")))
    ctx.stock_inicial(item, f, "und")
    ctx.items[item.nombre.lower()] = item
    ctx.nuevos.add(item.id)
    ctx.rep.platos += 1


def _recetas(ctx: _Contexto, filas: list[tuple[int, dict]]):
    """Agrupa las filas por plato; cada plato se guarda completo (o sus errores quedan en su primera fila)."""
    grupos: dict[str, list[tuple[int, dict]]] = {}
    for n, f in filas:
        grupos.setdefault(f["plato"].lower(), []).append((n, f))
    for _, grupo in grupos.items():
        n0, f0 = grupo[0]
        try:
            plato = ctx.items.get(f0["plato"].lower())
            if plato is None:
                raise ErrorNegocio(f"El plato '{f0['plato']}' no existe (ni en la planilla ni en el sistema).")
            if plato.id not in ctx.nuevos and ctx.db.query(Receta).filter_by(item_id=plato.id).first():
                raise ErrorNegocio(f"'{plato.nombre}' ya tiene receta: edítala desde el sistema.")
            lineas, filas_error = [], False
            for n, f in grupo:
                try:
                    ing = ctx.items.get(f["ingrediente"].lower())
                    if ing is None:
                        raise ErrorNegocio(f"El ingrediente '{f['ingrediente']}' no existe.")
                    if ing.tipo not in TIPOS_INGREDIENTE:
                        raise ErrorNegocio(f"'{ing.nombre}' no puede ser ingrediente.")
                    unidad = _unidad(ctx.db, f["unidad"]).codigo if f.get("unidad") else None
                    cantidad = _numero(f["cantidad"], "La cantidad", obligatorio=True)
                    lineas.append(LineaReceta(ing.id, cantidad, unidad, _si_no(f.get("removible", ""))))
                except (ErrorNegocio, ValueError) as e:
                    ctx.rep.errores.append(ErrorFila("recetas", n, str(e)))
                    filas_error = True
            if filas_error:
                continue
            rend = next(((f["rendimiento"], f.get("unidad_rendimiento") or None) for _, f in grupo if f.get("rendimiento")),
                        None)
            if rend is None and plato.tipo == TipoItem.componente:
                raise ErrorNegocio("Un componente necesita 'rendimiento' (cuánto sale de la receta).")
            guardar_receta(ctx.db, empresa_id=ctx.empresa_id, item_id=plato.id,
                           rendimiento=_numero(rend[0], "El rendimiento") if rend else 1,
                           rendimiento_unidad=rend[1] if rend else None, lineas=lineas)
            ctx.rep.recetas += 1
        except (ErrorNegocio, ValueError) as e:
            ctx.rep.errores.append(ErrorFila("recetas", n0, str(e)))


def procesar(db: Session, *, empresa_id: int, sucursal_id: Optional[int], contenido: bytes) -> Reporte:
    """
    Aplica la planilla dentro de la transacción actual. Quien llama decide: si `rep.valido` y quiere
    importar, hace commit; en cualquier otro caso, rollback (el servicio no confirma nada).
    """
    rep = Reporte()
    libro = abrir_libro(contenido)
    datos = {k: _leer_hoja(libro, k, rep) for k in HOJAS}
    if not rep.errores and not any(datos.values()):
        raise ErrorNegocio("La planilla está vacía: llena al menos una hoja (insumos, platos o recetas).")
    ctx = _Contexto(db, empresa_id, sucursal_id, rep)
    for clave, manejador in (("insumos", _insumo), ("platos", _plato)):
        for n, fila in datos[clave]:
            try:
                for c in OBLIGATORIAS[clave]:
                    if not fila.get(c):
                        raise ErrorNegocio(f"Falta {c}.")
                manejador(ctx, fila)
            except (ErrorNegocio, ValueError) as e:
                rep.errores.append(ErrorFila(clave, n, str(e)))
    filas_rec = []
    for n, f in datos["recetas"]:
        if not f.get("plato") or not f.get("ingrediente") or not f.get("cantidad"):
            rep.errores.append(ErrorFila("recetas", n, "Cada fila necesita plato, ingrediente y cantidad."))
        else:
            filas_rec.append((n, f))
    _recetas(ctx, filas_rec)
    rep.errores.sort(key=lambda e: (list(HOJAS).index(e.hoja) if e.hoja in HOJAS else 9, e.fila))
    return rep


# ── Plantilla ────────────────────────────────────────────────
def plantilla_xlsx() -> bytes:
    libro = Workbook()
    ayuda = libro.active
    ayuda.title = "Instrucciones"
    lineas = [
        "CARGA MASIVA — Yepar Restaurant",
        "",
        "1. Llena las hojas 'insumos', 'platos' y 'recetas' (puedes dejar alguna vacía). No cambies los títulos de las columnas.",
        "2. Todo se revisa antes de guardar: si hay un error no entra nada y se te dice la hoja y la fila.",
        "3. Los nombres no se pueden repetir ni existir ya en el sistema. Esta carga solo crea, no edita.",
        "",
        "INSUMOS — tipo: materia_prima o componente · unidad: g, kg, ml, l, und...",
        "   ejemplo: Tomate | materia_prima | kg | Verduras | 5 | 20 | 1500   (stock_minimo y stock_inicial en esa unidad; costo por esa unidad, sin IVA)",
        "PLATOS — tipo: preparado o reventa · estacion: cocina o barra (opcional)",
        "   ejemplo: Hamburguesa | preparado | Platos | 6500 | cocina",
        "   ejemplo: Cola lata | reventa | Bebidas | 1500 | barra | 24 | 700   (la reventa sí puede traer stock inicial)",
        "RECETAS — una fila por ingrediente. removible: si/no (el cliente puede pedirlo 'sin'). rendimiento solo en componentes.",
        "   ejemplo: Hamburguesa | Pan | 1 | und | no     y     Hamburguesa | Cebolla | 30 | g | si",
        "",
        "El stock inicial entra en la sucursal que elijas al subir el archivo.",
    ]
    for i, t in enumerate(lineas, start=1):
        ayuda.cell(row=i, column=1, value=t)
    ayuda["A1"].font = Font(bold=True, size=14)
    ayuda.column_dimensions["A"].width = 120
    relleno = PatternFill("solid", fgColor="E9E3FF")
    for clave, columnas in HOJAS.items():
        h = libro.create_sheet(clave)
        for c, nombre in enumerate(columnas, start=1):
            celda = h.cell(row=1, column=c, value=nombre)
            celda.font, celda.fill = Font(bold=True), relleno
            h.column_dimensions[celda.column_letter].width = 20
        h.freeze_panes = "A2"
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


# ── Asistente de arranque ────────────────────────────────────
def estado_arranque(db: Session, empresa_id: int) -> list[dict]:
    """La lista de pasos para dejar el restaurante andando, con lo que ya está hecho."""
    def n(modelo, *filtros):
        return db.query(func.count()).select_from(modelo).filter(*filtros).scalar() or 0

    def items(*tipos):
        return n(Item, Item.empresa_id == empresa_id, Item.activo.is_(True), Item.tipo.in_(tipos))

    insumos = items(TipoItem.materia_prima, TipoItem.componente)
    platos = items(TipoItem.vendible_preparado, TipoItem.vendible_reventa)
    conreceta = n(Receta, Receta.empresa_id == empresa_id)
    personal = (db.query(func.count(func.distinct(UsuarioRol.usuario_id))).join(Usuario, Usuario.id == UsuarioRol.usuario_id)
                  .filter(Usuario.empresa_id == empresa_id, UsuarioRol.rol != RolUsuario.administrador).scalar() or 0)
    con_stock = n(MovimientoStock, MovimientoStock.empresa_id == empresa_id,
                  MovimientoStock.tipo == TipoMovimientoStock.compra)
    pasos = [
        ("sucursal", "Crear tu sucursal", n(Sucursal, Sucursal.empresa_id == empresa_id) > 0, "Tu local ya está creado."),
        ("insumos", "Cargar insumos (ingredientes y bebidas)", insumos > 0, f"{insumos} cargados. Usa la planilla Excel."),
        ("platos", "Cargar la carta (platos y bebidas a la venta)", platos > 0, f"{platos} cargados con su precio."),
        ("recetas", "Armar las recetas de los platos", conreceta > 0, f"{conreceta} recetas listas."),
        ("stock", "Ingresar el stock inicial", con_stock > 0, "Sin stock el sistema avisará que todo falta."),
        ("mesas", "Dibujar las mesas del salón", n(Mesa, Mesa.empresa_id == empresa_id) > 0, "Crea las mesas con su plano."),
        ("personal", "Crear al personal (meseros, cocina, caja)", personal > 0, f"{personal} personas."),
        ("caja", "Crear la caja del local", n(Caja, Caja.empresa_id == empresa_id) > 0, "Necesaria para abrir turnos y cobrar."),
        ("cobro", "Definir cómo se cobra la propina", n(ConfiguracionCobro, ConfiguracionCobro.empresa_id == empresa_id) > 0,
         "Elige el modo y el porcentaje sugerido."),
        ("facturacion", "Conectar tu facturador (YeparDTE o YeparFactura)",
         n(IntegracionFacturacion, IntegracionFacturacion.empresa_id == empresa_id) > 0, "Para emitir boletas y facturas."),
    ]
    return [{"clave": k, "titulo": t, "hecho": h, "detalle": d} for k, t, h, d in pasos]
