# ============================================================
# YEPAR RESTAURANT — backend/tests/test_api_carga_masiva.py
# Carga masiva por Excel (todo o nada) y asistente de arranque.
# ============================================================

import io

import pytest
from openpyxl import Workbook, load_workbook

from models import RolUsuario

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
COL = {"insumos": ["nombre", "tipo", "unidad", "categoria", "stock_minimo", "stock_inicial", "costo_por_unidad"],
       "platos": ["nombre", "tipo", "categoria", "precio", "estacion", "stock_inicial", "costo_por_unidad"],
       "recetas": ["plato", "ingrediente", "cantidad", "unidad", "removible", "rendimiento", "unidad_rendimiento"]}


def libro(**hojas):
    """libro(insumos=[[...], ...]) → bytes de un .xlsx con los encabezados de la plantilla."""
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        h = wb.create_sheet(nombre)
        h.append(COL[nombre])
        for f in filas:
            h.append(f + [None] * (len(COL[nombre]) - len(f)))
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


@pytest.fixture()
def api(db, cliente, cabeceras, local, u, nuevo_usuario):
    empresa, sucursal = local
    personas = {
        "admin": nuevo_usuario(empresa, "Admin", roles=(RolUsuario.administrador,), sucursal_id=sucursal.id),
        "chef": nuevo_usuario(empresa, "Chef", roles=(RolUsuario.cocinero,), sucursal_id=sucursal.id),
    }

    class Api:
        def get(self, p, ruta):
            return cliente.get(ruta, headers=cabeceras(personas[p]))

        def post(self, p, ruta, json=None):
            return cliente.post(ruta, headers=cabeceras(personas[p]), json=json or {})

        def subir(self, p, ruta, contenido, nombre="carga.xlsx"):
            return cliente.post(ruta, headers=cabeceras(personas[p]), files={"archivo": (nombre, contenido, XLSX)})

    return Api()


PLANILLA = dict(
    insumos=[["Tomate", "materia_prima", "kg", "Verduras", "2", "10", "1500"],
             ["Pan", "materia_prima", "und", "Panadería", "10", "50", "300"],
             ["Carne", "materia_prima", "g", "Carnes", None, "3000", "9"],
             ["Salsa de la casa", "componente", "g", "Salsas"]],
    platos=[["Hamburguesa", "preparado", "Platos", "6500", "cocina"],
            ["Cola lata", "reventa", "Bebidas", "1500", "barra", "24", "700"]],
    recetas=[["Hamburguesa", "Pan", "1", "und", "no"], ["Hamburguesa", "Carne", "150", "g", "no"],
             ["Hamburguesa", "Tomate", "30", "g", "si"],
             ["Salsa de la casa", "Tomate", "1", "kg", "no", "1", "kg"]],
)


def stock(api):
    return {f["nombre"]: f["fisico"] for f in api.get("admin", "/inventario/stock").json()}


def test_plantilla_es_un_excel_con_las_hojas(api):
    r = api.get("admin", "/carga-masiva/plantilla")
    assert r.status_code == 200 and r.headers["content-type"].startswith(XLSX)
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Instrucciones", "insumos", "platos", "recetas"]
    assert [c.value for c in wb["platos"][1]] == COL["platos"]


def test_importar_crea_todo_con_stock_inicial_y_recetas(db, u, api):
    r = api.subir("admin", "/carga-masiva/importar", libro(**PLANILLA))
    assert r.status_code == 201, r.text
    rep = r.json()
    assert rep["aplicado"] and rep["errores"] == []
    assert (rep["insumos"], rep["platos"], rep["recetas"], rep["con_stock_inicial"]) == (4, 2, 2, 4)
    assert rep["categorias_nuevas"] == 6
    s = stock(api)
    assert s["Tomate"] == "10000.000" and s["Pan"] == "50.000" and s["Carne"] == "3000.000" and s["Cola lata"] == "24.000"
    tomate = next(i for i in api.get("admin", "/items").json() if i["nombre"] == "Tomate")
    assert tomate["unidad_base"] == "g" and tomate["stock_minimo"] == "2000.000"
    from decimal import Decimal
    from models import Item
    # El costo se pasó a unidad base: $1.500 el kilo → $1,5 el gramo.
    assert db.get(Item, tomate["id"]).costo_promedio == Decimal("1.5")
    hamburguesa = next(i for i in api.get("admin", "/items").json() if i["nombre"] == "Hamburguesa")
    assert api.get("admin", f"/items/{hamburguesa['id']}/receta").status_code == 200
    assert api.get("admin", "/carga-masiva/asistente").json()["completados"] >= 4


def test_validar_no_guarda_nada(db, u, api):
    r = api.subir("admin", "/carga-masiva/validar", libro(**PLANILLA))
    assert r.status_code == 200 and r.json()["valido"] and r.json()["aplicado"] is False
    assert api.get("admin", "/items").json() == []


def test_un_error_y_no_entra_nada_con_fila_y_motivo(db, u, api):
    malo = {**PLANILLA, "platos": [["Hamburguesa", "preparado", "Platos", "abc"], ["Cola", "raro", "", "1500"]]}
    r = api.subir("admin", "/carga-masiva/importar", libro(**malo))
    assert r.status_code == 422
    cuerpo = r.json()
    assert "no se guardó nada" in cuerpo["detalle"]
    errores = {(e["hoja"], e["fila"]): e["mensaje"] for e in cuerpo["errores"]}
    assert "no es un número" in errores[("platos", 2)] and "no válido" in errores[("platos", 3)]
    assert any(h == "recetas" for h, _ in errores)            # la receta apunta a un plato que no se creó
    assert api.get("admin", "/items").json() == []             # ni siquiera los insumos buenos quedaron
    assert api.get("admin", "/categorias").json() == []


def test_reglas_de_cada_fila(db, u, api):
    r = api.subir("admin", "/carga-masiva/validar", libro(
        insumos=[["Tomate", "materia_prima", "kg", None, None, "5"],              # stock sin costo
                 ["Sal", "materia_prima", "xx"],                                   # unidad inexistente
                 ["Salsa", "componente", "g", None, None, "5", "1"],               # componente con stock
                 ["Arroz", "materia_prima", "kg"], ["arroz", "materia_prima", "kg"],  # repetido
                 ["", "materia_prima", "kg"]],                                      # sin nombre
        platos=[["Pasta", "preparado", None, "5000", "cocina", "3", "1"]],         # preparado con stock
        recetas=[["Pasta", "Arroz", "1", "kg", "quizas"], ["Pasta", "Fantasma", "1"]]))
    errores = {(e["hoja"], e["fila"]): e["mensaje"] for e in r.json()["errores"]}
    assert "costo por unidad" in errores[("insumos", 2)]
    assert "no existe" in errores[("insumos", 3)]
    assert "no lleva stock" in errores[("insumos", 4)]
    assert "Ya existe" in errores[("insumos", 6)]
    assert "Falta nombre" in errores[("insumos", 7)]
    assert "no tiene stock propio" in errores[("platos", 2)]
    assert "no existe" in errores[("recetas", 2)]              # el plato con error no se creó: su receta se avisa una vez


def test_no_pisa_lo_que_ya_existe_y_la_receta_ajena_no_se_reemplaza(db, u, api):
    assert api.subir("admin", "/carga-masiva/importar", libro(**PLANILLA)).status_code == 201
    r = api.subir("admin", "/carga-masiva/validar", libro(
        insumos=[["Tomate", "materia_prima", "kg"]],
        recetas=[["Hamburguesa", "Pan", "2", "und"]]))
    errores = {(e["hoja"], e["fila"]): e["mensaje"] for e in r.json()["errores"]}
    assert "Ya existe" in errores[("insumos", 2)] and "ya tiene receta" in errores[("recetas", 2)]


def test_archivo_invalido_vacio_y_permisos(db, u, api):
    assert api.subir("admin", "/carga-masiva/importar", b"no soy excel").status_code == 422
    assert api.subir("admin", "/carga-masiva/importar", libro(insumos=[["x", "materia_prima", "g"]]), "datos.csv").status_code == 422
    assert api.subir("admin", "/carga-masiva/importar", libro(insumos=[], platos=[])).status_code == 422
    sin_col = Workbook()
    sin_col.active.title = "insumos"
    sin_col.active.append(["nombre"])
    b = io.BytesIO(); sin_col.save(b)
    r = api.subir("admin", "/carga-masiva/validar", b.getvalue())
    assert "Falta la columna" in r.json()["errores"][0]["mensaje"]
    assert api.subir("chef", "/carga-masiva/importar", libro(**PLANILLA)).status_code == 403
    assert api.get("chef", "/carga-masiva/plantilla").status_code == 403
    assert api.get("chef", "/carga-masiva/asistente").status_code == 403


def test_asistente_marca_lo_hecho_y_el_siguiente_paso(db, u, api):
    a = api.get("admin", "/carga-masiva/asistente").json()
    hechos = {p["clave"]: p["hecho"] for p in a["pasos"]}
    assert hechos["sucursal"] and not hechos["insumos"] and not hechos["mesas"]
    assert a["siguiente"] == "Cargar insumos (ingredientes y bebidas)" and a["total"] == 10
    api.subir("admin", "/carga-masiva/importar", libro(**PLANILLA))
    a = api.get("admin", "/carga-masiva/asistente").json()
    hechos = {p["clave"]: p["hecho"] for p in a["pasos"]}
    assert hechos["insumos"] and hechos["platos"] and hechos["recetas"] and hechos["stock"] and not hechos["caja"]
    assert a["siguiente"] == "Dibujar las mesas del salón"
